"""Training-only content acquisition. No human evidence or model scores are read.

Private payloads never enter repository files, Actions artifacts, or job logs.
Run-specific objects are immutable. The gateway verifies SHA after upload/download.
"""
import base64
import datetime as dt
import hashlib
import io
import json
import os
import re
import subprocess
import tempfile
import time
import unicodedata
import zipfile
from pathlib import Path
from urllib.parse import urljoin, urlparse, unquote

import requests
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright

GATEWAY = 'https://sshovpnepgswzvjwjuyz.supabase.co/functions/v1/autonomous-corpus-gateway'
AUDIENCE = 'cy-school-news-autonomous-training'
MAX_BYTES = 10_000_000
AS_OF = '2026-10-01'
METHOD_VERSION = 'autonomous-content-v1'
WORK = Path(tempfile.mkdtemp(prefix='private-corpus-'))
RUN = os.environ['GITHUB_RUN_ID']
ROOT = f'runs/{RUN}/'
session = requests.Session()
session.headers['User-Agent'] = 'CYSchoolNews-Research/1.0 (bounded public-source corpus acquisition)'
session.headers['Accept'] = '*/*'

def digest(b): return hashlib.sha256(b).hexdigest()
def canonical(x): return json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()
def clean(s): return re.sub(r'[ \t\u3000]+',' ',unicodedata.normalize('NFC',str(s))).replace('\r','').strip()
def meaningful(text,title=''):
    text = text.replace(title,'')
    for noise in ['瀏覽數','最後更新日期','發布日期','友善列印','分享','附件下載','回上一頁']:
        text=text.replace(noise,'')
    return len(re.findall(r'[\w\u4e00-\u9fff]',text)) >= 80 and len(set(text))>=30

def token():
    r=session.get(os.environ['ACTIONS_ID_TOKEN_REQUEST_URL']+'&audience='+AUDIENCE,
                  headers={'Authorization':'Bearer '+os.environ['ACTIONS_ID_TOKEN_REQUEST_TOKEN']},timeout=20)
    r.raise_for_status();return r.json()['value']

def put(path,data):
    h=digest(data)
    r=session.post(GATEWAY,headers={'Authorization':'Bearer '+token()},
        json={'op':'put','path':ROOT+path,'data':base64.b64encode(data).decode(),'sha256':h},timeout=60)
    if not r.ok: raise RuntimeError(f'private upload failed status={r.status_code} code={r.json().get("error")}')
    assert r.json()['verified'] and r.json()['sha256']==h
    (WORK/path).parent.mkdir(parents=True,exist_ok=True);(WORK/path).write_bytes(data)
    return {'path':ROOT+path,'sha256':h,'size':len(data),'readback_verified':True}

def official(url,school):
    p=urlparse(url);host=p.hostname or ''
    return p.scheme=='https' and not p.username and not p.password and (host==f'{school}.cy.edu.tw' or host.endswith(f'.{school}.cy.edu.tw'))

def fetch(url,school,context=None):
    if not official(url,school):raise ValueError('non-official URL')
    for attempt in range(2):
        try:
            r=session.get(url,timeout=(12,35),stream=True,allow_redirects=False)
            for _ in range(4):
                if r.is_redirect:
                    target=urljoin(r.url,r.headers['Location'])
                    if not official(target,school):raise ValueError('external redirect')
                    r=session.get(target,timeout=(12,35),stream=True,allow_redirects=False)
                else:break
            r.raise_for_status();data=bytearray()
            for chunk in r.iter_content(65536):
                data.extend(chunk)
                if len(data)>MAX_BYTES:raise ValueError('size limit')
            return bytes(data),{'method':'direct_get','status':r.status_code,'final_url':r.url,'content_type':r.headers.get('Content-Type','')}
        except Exception:
            if attempt==0:time.sleep(1)
    if context:
        r=context.request.get(url,timeout=40000)
        if r.ok and official(r.url,school):
            b=r.body()
            if len(b)<=MAX_BYTES:return b,{'method':'browser_session_request','status':r.status,'final_url':r.url,'content_type':r.headers.get('content-type','')}
    raise RuntimeError('bounded download failed')

def body_extract(raw,title):
    soup=BeautifulSoup(raw.decode('utf-8','replace'),'html.parser')
    for n in soup.select('script,style,nav,header,footer,form,.breadcrumb,.share,.social,.mptattach'):
        n.decompose()
    node=None
    for sel in ['.meditor','.mpgdetail','.news_content','.news-content','article','#Dyn_2_2']:
        candidates=soup.select(sel)
        if candidates:
            node=max(candidates,key=lambda n:len(n.get_text()));break
    if node is None:return []
    blocks=[]
    for n in node.find_all(['p','li','tr','h2','h3','h4']):
        if n.find_parent(['p','li','tr']) is not None:continue
        t=clean(n.get_text(' ',strip=True))
        if t and t!=title and t not in [x['text'] for x in blocks]:
            blocks.append({'locator':f'body:{len(blocks)+1}','text':t})
    if not blocks:
        t=clean(node.get_text('\n',strip=True))
        if t and t!=title:blocks=[{'locator':'body:1','text':t}]
    return blocks

def discover(raw,url,school):
    soup=BeautifulSoup(raw.decode('utf-8','replace'),'html.parser');out=[];seen=set()
    for a in soup.select('a[href]'):
        link=urljoin(url,a['href']);name=clean(a.get_text(' ',strip=True));path=unquote(urlparse(link).path)
        m=re.search(r'\.(pdf|docx?|xlsx?|pptx|png|jpe?g)(?:\b|$)',name+' '+path,re.I)
        if not m:continue
        if not official(link,school) or link in seen:continue
        seen.add(link);ext=m.group(1).lower();filename=name or path.rsplit('/',1)[-1]
        out.append({'url':link,'filename':filename,'extension':ext})
    return out

def parse_attachment(data,ext):
    units=[];method=ext
    if ext in ['docx','xlsx','pptx']:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist())>40_000_000:raise ValueError('expanded size limit')
    if ext=='pdf':
        from pypdf import PdfReader
        pages=PdfReader(io.BytesIO(data)).pages
        if len(pages)>60:raise ValueError('page limit')
        for i,p in enumerate(pages):
            text=clean(p.extract_text() or '');parse_method='pdf_text'
            if len(re.findall(r'\w',text))<30:
                with tempfile.TemporaryDirectory() as tmp:
                    pdf=Path(tmp)/'source.pdf';pdf.write_bytes(data);png=Path(tmp)/'page'
                    subprocess.run(['pdftoppm','-f',str(i+1),'-l',str(i+1),'-scale-to','2200','-singlefile','-png',str(pdf),str(png)],capture_output=True,timeout=30,check=True)
                    o=subprocess.run(['tesseract',str(png)+'.png','stdout','-l','chi_tra+eng','--psm','3'],capture_output=True,timeout=40,check=True)
                    text=clean(o.stdout.decode('utf-8'));parse_method='pdf_image_ocr'
            if text:units.append({'locator':f'page:{i+1}','page':i+1,'text':text,'parse_method':parse_method})
    elif ext=='docx':
        from docx import Document
        d=Document(io.BytesIO(data))
        for i,p in enumerate(d.paragraphs):
            if clean(p.text):units.append({'locator':f'paragraph:{i+1}','text':clean(p.text)})
        for i,t in enumerate(d.tables):
            text='\n'.join(' | '.join(clean(c.text) for c in row.cells) for row in t.rows)
            if text:units.append({'locator':f'table:{i+1}','text':text})
    elif ext=='xlsx':
        from openpyxl import load_workbook
        w=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        for s in w:
            if s.max_row>20000 or s.max_column>100:raise ValueError('sheet size limit')
            lines=[]
            for i,row in enumerate(s.iter_rows(values_only=True)):
                vals=[v.isoformat() if isinstance(v,(dt.datetime,dt.date)) else str(v) if v is not None else '' for v in row]
                if any(vals):lines.append(f'row {i+1}: '+ ' | '.join(vals))
            if lines:units.append({'locator':f'sheet:{s.title}','sheet':s.title,'text':'\n'.join(lines)})
    elif ext=='pptx':
        from pptx import Presentation
        p=Presentation(io.BytesIO(data))
        for i,s in enumerate(p.slides):
            text='\n'.join(clean(shape.text) for shape in s.shapes if shape.has_text_frame)
            if text:units.append({'locator':f'slide:{i+1}','slide':i+1,'text':text})
    elif ext=='xls':
        import xlrd
        w=xlrd.open_workbook(file_contents=data)
        for s in w.sheets():
            if s.nrows>20000 or s.ncols>100:raise ValueError('sheet size limit')
            text='\n'.join(f'row {i+1}: '+' | '.join(str(v) for v in s.row_values(i)) for i in range(s.nrows))
            if text:units.append({'locator':f'sheet:{s.name}','sheet':s.name,'text':clean(text)})
    elif ext=='doc':
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'source.doc';path.write_bytes(data)
            r=subprocess.run(['antiword','-mUTF-8',str(path)],capture_output=True,timeout=20,check=True)
            text=clean(r.stdout.decode('utf-8'));units=[{'locator':'document:1','text':text}] if text else []
    elif ext in ['png','jpg','jpeg']:
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/('image.'+ext);path.write_bytes(data)
            r=subprocess.run(['tesseract',str(path),'stdout','-l','chi_tra+eng','--psm','3'],capture_output=True,timeout=40,check=True)
            text=clean(r.stdout.decode('utf-8'));units=[{'locator':'image:1','text':text,'parse_method':'image_ocr'}] if text else []
    else:raise ValueError('parser unavailable')
    return {'parse_method':method,'units':units,'meaningful':meaningful('\n'.join(x['text'] for x in units))}

def acquire(item,context):
    school=item['school'];aid=item['id'];url=item['url'];title=item['title']
    meta={k:item.get(k) for k in ['id','school','title','date','category','source_category','url','first_seen']}
    meta['captured_at']=dt.datetime.now(dt.timezone.utc).isoformat();meta['parser_version']=METHOD_VERSION
    page=None;raw=None;acq=None
    try:raw,acq=fetch(url,school,context)
    except Exception:pass
    blocks=body_extract(raw,title) if raw else []
    attachments=discover(raw,url,school) if raw else []
    if not raw or not meaningful('\n'.join(x['text'] for x in blocks),title):
        try:
            page=context.new_page();resp=page.goto(url,wait_until='domcontentloaded',timeout=45000);page.wait_for_timeout(1800)
            rendered=page.content().encode();rendered_blocks=body_extract(rendered,title)
            if raw:meta['direct_raw']=put(f'{aid}/direct.html',raw)
            if rendered_blocks or not raw:
                raw=rendered;blocks=rendered_blocks;attachments=discover(raw,url,school)
                acq={'method':'playwright_rendered_dom','status':resp.status if resp else None,'final_url':page.url,'content_type':'text/html'}
        except Exception:meta['browser_fallback']='failed'
    if not raw:meta['gate']='FAIL';meta['failure']='detail_unavailable';return meta
    meta['raw']=put(f'{aid}/source.html',raw);meta['acquisition']=acq
    meta['body']=put(f'{aid}/body.json',canonical(blocks));meta['body_meaningful']=meaningful('\n'.join(x['text'] for x in blocks),title)
    parsed=[]
    for index,a in enumerate(attachments):
        entry={**a,'index':index,'status':'pending'}
        if index>=6:entry['status']='resource_limit';parsed.append(entry);continue
        try:
            data,m=fetch(a['url'],school,context);entry['acquisition']=m
            # The URL came from current DOM. A browser-session click is the next fallback.
        except Exception:
            try:
                if page is None:
                    page=context.new_page();page.goto(url,wait_until='domcontentloaded',timeout=40000)
                with page.expect_download(timeout=20000) as info:
                    page.locator('a').filter(has_text=a['filename']).first.click()
                dl=info.value;data=Path(dl.path()).read_bytes();entry['acquisition']={'method':'browser_session_download'}
                if len(data)>MAX_BYTES:raise ValueError('size limit')
            except Exception:entry['status']='download_failed';parsed.append(entry);continue
        entry['raw']=put(f'{aid}/attachment-{index}.{a["extension"]}',data)
        try:
            extracted=parse_attachment(data,a['extension'])
            entry['extracted']=put(f'{aid}/attachment-{index}.json',canonical(extracted));entry['meaningful']=extracted['meaningful'];entry['parse_method']=extracted['parse_method']
            entry['status']='parsed' if extracted['meaningful'] else 'not_meaningful'
        except Exception as e:entry['status']='parse_failed';entry['error_type']=type(e).__name__
        parsed.append(entry)
    if page:page.close()
    meta['attachments']=parsed;meta['attachment_meaningful']=any(a.get('meaningful') for a in parsed)
    meta['gate']='PASS' if meta['body_meaningful'] or meta['attachment_meaningful'] else 'FAIL'
    # Offline reconstruction must use the captured artifact, never a fresh website request.
    meta['body_reproducible']=digest(canonical(body_extract(raw,title)))==meta['body']['sha256']
    for a in parsed:
        if a.get('extracted'):
            source=WORK/a['raw']['path'].removeprefix(ROOT)
            try:a['reproducible']=digest(canonical(parse_attachment(source.read_bytes(),a['extension'])))==a['extracted']['sha256']
            except Exception:a['reproducible']=False
    meta['reproducible']=meta['body_reproducible'] and all(a.get('reproducible',True) for a in parsed)
    if not meta['reproducible']:meta['gate']='FAIL'
    put(f'{aid}/record.json',canonical(meta))
    print(json.dumps({'id':aid,'gate':meta['gate'],'body':meta['body_meaningful'],'attachments':len(parsed),'parsed':sum(a['status']=='parsed' for a in parsed),'reproducible':meta['reproducible']}),flush=True)
    time.sleep(1)
    return meta

def pool_select(items):
    pool=[]
    today=dt.date.fromisoformat(AS_OF)
    for school in ['cysh','cygsh']:
        groups={}
        for item in items:
            if item.get('school')!=school or not official(item.get('url',''),school):continue
            try:age=(today-dt.date.fromisoformat(item['date'][:10])).days
            except Exception:continue
            if age<0:continue
            agebin=0 if age<=7 else 1 if age<=30 else 2 if age<=90 else 3 if age<=365 else 4
            strat=(agebin,item.get('category',''),bool(item.get('attachments') or item.get('has_attachments')))
            groups.setdefault(strat,[]).append(item)
        for group in groups.values():group.sort(key=lambda x:(x['date'],x['id']),reverse=True)
        selected=[]
        while len(selected)<30:
            added=False
            for key in sorted(groups):
                if groups[key] and len(selected)<30:selected.append(groups[key].pop(0));added=True
            if not added:break
        pool.extend(selected)
    assert len(pool)==60 and len({x['id'] for x in pool})==60
    return pool

def freeze(records):
    chosen=[]
    for school in ['cysh','cygsh']:
        passed=[x for x in records if x['school']==school and x['gate']=='PASS']
        # Selection uses acquisition coverage only, never semantic inference or labels.
        rich=sorted([x for x in passed if x.get('attachment_meaningful')],key=lambda x:x['id'])
        only=sorted([x for x in passed if not x.get('attachments') and x.get('body_meaningful')],key=lambda x:x['id'])
        first=rich[:5]+only[:3]
        remainder=sorted([x for x in passed if x['id'] not in {y['id'] for y in first}],key=lambda x:x['id'])
        selected=(first+remainder)[:15]
        if len(selected)<15:return None
        chosen.extend(selected)
    if sum(not x.get('attachments') and x['body_meaningful'] for x in chosen)<5 or sum(x['attachment_meaningful'] for x in chosen)<5:return None
    return chosen

def main():
    source=session.get('https://raw.githubusercontent.com/tsaibohau/cy-school-news/main/docs/data/announcements.json',timeout=30);source.raise_for_status()
    index=source.json();pool=pool_select(index['items'])
    put('source-index.json',source.content)
    put('candidate-pool.json',canonical(pool))
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=True);context=browser.new_context(accept_downloads=True)
        # First candidate from each school is a smoke, then proceed without a human pause.
        smoke=[next(x for x in pool if x['school']==s) for s in ['cysh','cygsh']]
        order=smoke+[x for x in pool if x['id'] not in {v['id'] for v in smoke}]
        records=[]
        for item in order:
            try:record=acquire(item,context)
            except Exception as e:
                record={'id':item['id'],'school':item['school'],'gate':'FAIL','failure':type(e).__name__}
                print(json.dumps(record),flush=True)
            records.append(record)
            put(f'progress-{len(records):02d}.json',canonical(records))
        browser.close()
    chosen=freeze(records)
    report={'run_id':RUN,'commit':os.environ['GITHUB_SHA'],'as_of':AS_OF,'candidate_pool':len(pool),
        'pass_pool':sum(x['gate']=='PASS' for x in records),'frozen_count':len(chosen or []),
        'schools':{s:sum(x['school']==s for x in chosen or []) for s in ['cysh','cygsh']},
        'body_coverage':sum(x.get('body_meaningful',False) for x in chosen or []),
        'attachment_coverage':sum(x.get('attachment_meaningful',False) for x in chosen or []),
        'attachment_discovered':sum(len(x.get('attachments',[])) for x in records),
        'attachment_parsed':sum(a.get('status')=='parsed' for x in records for a in x.get('attachments',[])),
        'reproducibility':all(x.get('reproducible') for x in chosen) if chosen else False,
        'status':'CORPUS_FROZEN' if chosen else 'POOL_EXTENSION_REQUIRED'}
    put('acquisition-records.json',canonical(records))
    if chosen:
        payload=[]
        for record in chosen:
            body=json.loads((WORK/record['body']['path'].removeprefix(ROOT)).read_bytes())
            atts=[]
            for a in record['attachments']:
                extracted=json.loads((WORK/a['extracted']['path'].removeprefix(ROOT)).read_bytes()) if a.get('extracted') else None
                atts.append({**a,'content':extracted})
            payload.append({**record,'body_content':body,'attachment_content':atts})
        frozen={'version':'v3.4-pilot','as_of':AS_OF,'persona':{'grade':1,'role':'student','school':'record.school','interests':[],'read_state':'unread','prior_actions':'unknown'},'records':payload}
        # Persona/date are frozen before inference. No primitive or human truth is added.
        report['frozen']=put('frozen-corpus.json',canonical(frozen))
        report['frozen_manifest_hash']=digest(canonical(chosen))
        put('frozen-manifest.json',canonical(chosen))
    put('report.json',canonical(report))
    print(json.dumps(report),flush=True)

if __name__=='__main__':main()
